"""Bounded, local-only roster extraction. No macros, formulas or AI calls."""
import csv
import hashlib
import io
import itertools
import re
import unicodedata
import zipfile

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 100
ALIASES = {
    'name': ('姓名', '名字', '真实姓名', '会员姓名', '成员姓名', '人员姓名', '学生姓名', '用户姓名', '会员名称', '联系人', 'name', 'fullname', 'realname'),
    'phone': ('手机号', '手机号码', '手机', '联系电话', '联系手机', '电话号码', '电话', '联系方式', 'phone', 'mobile', 'telephone', 'phonenumber', 'mobilephone'),
    'gender': ('性别', '男女', 'gender', 'sex'),
    'level': ('自评级别', '自评等级', '羽毛球等级', '技术等级', '运动等级', '等级', '级别', 'level', 'rating'),
    'groups': ('组别', '分组', '所属组别', '所属分组', '小组', '组', '班级', 'groups', 'group', 'team'),
}


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header_key(value):
    text = unicodedata.normalize('NFKC', cell_text(value)).lower()
    text = re.sub(r'\([^)]*\)|（[^）]*）', '', text)
    return re.sub(r'[\s_\-*:：/]', '', text)


HEADERS = {header_key(alias): key for key, aliases in ALIASES.items() for alias in aliases}


def _sheets(content, filename):
    suffix = filename.rsplit('.', 1)[-1].lower()
    if suffix == 'csv':
        try:
            text, encoding = content.decode('utf-8-sig'), 'UTF-8'
        except UnicodeDecodeError:
            try:
                text, encoding = content.decode('gb18030'), 'GB18030'
            except UnicodeDecodeError:
                raise ValueError('CSV编码无效，请另存为UTF-8') from None
        if '\x00' in text:
            raise ValueError('CSV内容无效')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = 'excel'
        yield 'CSV', csv.reader(io.StringIO(text, newline=''), dialect, strict=True), encoding
    elif suffix == 'xlsx':
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 500 or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Excel解压后过大，请精简表格')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except ValueError:
            raise
        except Exception:
            raise ValueError('无法读取Excel，请使用未加密的.xlsx文件') from None
        try:
            if len(book.worksheets) > 20:
                raise ValueError('Excel最多20个工作表')
            for sheet in book.worksheets:
                if sheet.sheet_state != 'visible':
                    continue
                sheet.reset_dimensions()
                yield sheet.title, sheet.iter_rows(max_col=MAX_COLUMNS + 1, values_only=True), 'Excel (.xlsx)'
        finally:
            book.close()
    elif suffix == 'xls':
        import xlrd
        try:
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
        except Exception:
            raise ValueError('无法读取Excel，请使用未加密的.xls文件') from None
        try:
            if book.nsheets > 20:
                raise ValueError('Excel最多20个工作表')
            for sheet in book.sheets():
                if sheet.visibility:
                    continue
                yield sheet.name, (sheet.row_values(i) for i in range(sheet.nrows)), 'Excel (.xls)'
        finally:
            book.release_resources()
    else:
        raise ValueError('请上传.xlsx、.xls或.csv文件')


def extract(content, filename):
    if not isinstance(content, bytes) or not content or len(content) > MAX_BYTES:
        raise ValueError('名册文件须为1字节至2MiB')
    rows, mappings, skipped, encoding = [], [], [], ''
    source = hashlib.sha256(content).hexdigest()
    try:
        for sheet, values, encoding in _sheets(content, filename):
            reader = iter(values)
            lead = list(itertools.islice(reader, 20))
            candidates = []
            for index, raw in enumerate(lead):
                names = [HEADERS.get(header_key(v)) for v in raw]
                if {'name', 'phone'} & set(names):
                    candidates.append((len(set(names) - {None}), -index, raw, names))
            if not candidates:
                skipped.append(sheet)
                continue
            _, negative, header, names = max(candidates, key=lambda c: c[:2])
            offset = -negative
            recognized = [n for n in names if n]
            if len(set(recognized)) != len(recognized):
                raise ValueError(f'{sheet} 第{offset + 1}行存在多个同类列，请保留一个姓名/手机号等字段列')
            if any(cell_text(v) for v in header[MAX_COLUMNS:]):
                raise ValueError('表格最多100列')
            columns = {key: cell_text(header[i]) for i, key in enumerate(names) if key}
            mappings.append({'sheet': sheet, 'header_row': offset + 1, 'columns': columns,
                             'ignored': [cell_text(v) for v, key in zip(header, names) if cell_text(v) and not key]})
            for number, raw in enumerate(itertools.chain(lead[offset + 1:], reader), offset + 2):
                if number > MAX_ROWS + 21:
                    raise ValueError('每张表最多1000条记录，请删除多余空行或尾部内容')
                raw = [cell_text(v) for v in raw]
                if not any(raw):
                    continue
                if any(raw[MAX_COLUMNS:]):
                    raise ValueError('表格最多100列')
                item = {key: raw[i] if i < len(raw) else '' for i, key in enumerate(names) if key}
                if not any(item.values()):
                    continue
                if len(rows) >= MAX_ROWS:
                    raise ValueError('每次最多导入1000条记录')
                item.update(line=number, sheet=sheet, source_key=hashlib.sha256(f'{source}:{sheet}:{number}'.encode()).hexdigest())
                rows.append(item)
    except csv.Error:
        raise ValueError('CSV引号或行结构无效') from None
    if not rows:
        raise ValueError('未找到会员记录，请提供姓名或手机号列')
    return rows, {'encoding': encoding, 'mappings': mappings, 'skipped_sheets': skipped}
